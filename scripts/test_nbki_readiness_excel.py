"""Readiness workbook protects operator-visible data and Excel cells."""

import unittest
from io import BytesIO

from openpyxl import load_workbook

from backend.nbki_excel import make_readiness_xlsx


class ReadinessExcelTests(unittest.TestCase):
    def test_export_has_only_supplied_rows_and_literal_text(self):
        report = {'rows': [{'loan_id': 5, 'contract_no': '=1+1', 'status': 'active',
                            'uid': None, 'crm_issues': ['Нет УИД'],
                            'missing_fields': ['Подтверждённый УИД договора'],
                            'xml_ready': False}]}
        sheet = load_workbook(BytesIO(make_readiness_xlsx(report))).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['B2'].value, '=1+1')
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['G2'].value, 'Нет')


if __name__ == '__main__':
    unittest.main()
