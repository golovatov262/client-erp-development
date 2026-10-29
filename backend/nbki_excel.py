"""Admin-only XLSX reconciliation workbook for unresolved NBKI deals."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = (
    'Источник', 'Номер договора в источнике', 'УИД', 'Заемщик в источнике',
    'Дата договора в источнике', 'Сумма в источнике', 'Причина расхождения',
    'ID договора CRM', 'Номер договора CRM', 'Заемщик CRM',
    'Дата договора CRM', 'Сумма CRM', 'Основание приема',
    'Решение после проверки', 'Правильный номер договора CRM', 'Комментарий',
)


def _read_rows(cur, org_id: int, status: str, source: str) -> list[tuple]:
    if source == 'XML':
        cur.execute("""SELECT d.contract_no,d.uid,d.subject_name,d.deal_date,
            d.original_amount,d.match_reason,l.id,l.contract_no,
            CASE WHEN m.member_type='FL' THEN CONCAT_WS(' ',m.last_name,m.first_name,m.middle_name)
                 ELSE m.company_name END,l.start_date,l.amount,d.reported_accepted
            FROM nbki_deals d
            LEFT JOIN loans l ON l.org_id=d.org_id AND l.contract_no=d.contract_no
            LEFT JOIN members m ON m.id=l.member_id
            WHERE d.org_id=%s AND d.match_status=%s ORDER BY d.contract_no,d.uid""",
            (org_id, status))
    else:
        cur.execute("""SELECT d.contract_no,d.uid,d.subject_name,d.deal_date,
            d.original_amount,d.match_reason,l.id,l.contract_no,
            CASE WHEN m.member_type='FL' THEN CONCAT_WS(' ',m.last_name,m.first_name,m.middle_name)
                 ELSE m.company_name END,l.start_date,l.amount,FALSE
            FROM nbki_event_deals d
            LEFT JOIN loans l ON l.org_id=d.org_id AND l.contract_no=d.contract_no
            LEFT JOIN members m ON m.id=l.member_id
            WHERE d.org_id=%s AND d.match_status=%s AND d.role='borrower'
            ORDER BY d.contract_no,d.uid""", (org_id, status))
    result = []
    for row in cur.fetchall():
        *fields, reported_accepted = row
        basis = ('Квитанция ticket2, подпись не проверена' if reported_accepted else
                 'Нет квитанции') if source == 'XML' else 'Принято по указанию, не подтверждено НБКИ'
        result.append((source, *fields, basis, None, None, None))
    return result


def make_registry_xlsx(cur, org_id: int, status: str) -> bytes:
    if status not in ('review', 'not_in_crm'):
        raise ValueError('Экспорт доступен только для проверки и отсутствующих договоров')
    book = Workbook()
    book.remove(book.active)
    for source in ('XML', 'MXL'):
        sheet = book.create_sheet(source)
        sheet.append(HEADERS)
        for row in _read_rows(cur, org_id, status, source):
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24476A')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Prevent Excel formula execution from imported text.
                elif isinstance(cell.value, (date, datetime)):
                    cell.number_format = 'DD.MM.YYYY'
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate((12, 27, 45, 34, 18, 20, 48, 16, 27,
                                       34, 18, 20, 42, 28, 30, 40), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    book.save(output)
    return output.getvalue()


MISSING_UID_HEADERS = (
    'ID договора CRM', 'Номер договора CRM', 'Заемщик CRM', 'Статус займа',
    'Дата договора', 'Сумма договора', 'Кандидаты УИД из XML',
    'Кандидаты УИД из MXL', 'Решение после проверки', 'Подтвержденный УИД',
    'Комментарий менеджера',
)


def make_missing_uid_xlsx(cur, org_id: int) -> bytes:
    """All non-closed CRM loans without a safely linked borrower UID."""
    cur.execute("""SELECT l.id,l.contract_no,
        CASE WHEN m.member_type='FL' THEN CONCAT_WS(' ',m.last_name,m.first_name,m.middle_name)
             ELSE m.company_name END AS subject_name,
        l.status,l.start_date,l.amount,
        COALESCE((SELECT STRING_AGG(d.uid || ' [' || d.match_status ||
            COALESCE(': ' || d.match_reason,'') || ']', '; ' ORDER BY d.uid)
            FROM nbki_deals d WHERE d.org_id=l.org_id AND d.contract_no=l.contract_no
            AND d.deal_role='borrower'),'') AS xml_candidates,
        COALESCE((SELECT STRING_AGG(d.uid || ' [' || d.match_status ||
            COALESCE(': ' || d.match_reason,'') || ']', '; ' ORDER BY d.uid)
            FROM nbki_event_deals d WHERE d.org_id=l.org_id AND d.contract_no=l.contract_no
            AND d.role='borrower'),'') AS mxl_candidates
        FROM loans l JOIN members m ON m.id=l.member_id
        WHERE l.org_id=%s AND l.status IN ('active','overdue','holiday')
        AND NOT EXISTS (SELECT 1 FROM nbki_deals d
            WHERE d.loan_id=l.id AND d.match_status='matched')
        AND NOT EXISTS (SELECT 1 FROM nbki_event_deals d
            WHERE d.loan_id=l.id AND d.role='borrower' AND d.match_status='matched')
        ORDER BY l.status,l.contract_no""", (org_id,))
    book = Workbook()
    sheet = book.active
    sheet.title = 'Без УИД'
    sheet.append(MISSING_UID_HEADERS)
    for row in cur.fetchall():
        sheet.append((*row, None, None, None))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476A')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            elif isinstance(cell.value, (date, datetime)):
                cell.number_format = 'DD.MM.YYYY'
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((17, 29, 36, 18, 19, 21, 80, 80, 30, 45, 50), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def make_readiness_xlsx(report: dict) -> bytes:
    """Export a snapshot of current, non-closed loan data gaps for review."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Готовность НБКИ'
    sheet.append(('ID договора CRM', 'Номер договора', 'Статус займа', 'УИД',
                  'Замечания CRM', 'Блокирующие пункты для XML', 'Готовность XML'))
    for row in report['rows']:
        sheet.append((row['loan_id'], row['contract_no'], row['status'], row['uid'],
                      '; '.join(row['crm_issues']), '; '.join(row['missing_fields']),
                      'Нет' if not row['xml_ready'] else 'Да'))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476A')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((18, 30, 17, 46, 80, 110, 20), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def make_1c_reconciliation_xlsx(report: dict) -> bytes:
    """Export review signals with source date/row and untouched CRM values."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Сверка 1С и CRM'
    sheet.append(('ID CRM', 'Договор', 'УИД', 'Дата события 1С', 'Строка 1С',
                  'Показатель', 'Дата платежа', 'Значение 1С', 'Значение CRM',
                  'Результат', 'Кандидаты платежей CRM', 'Пояснение',
                  'Решение менеджера', 'Комментарий'))
    labels = {'contract_end_date': 'Дата окончания договора',
              'last_payment': 'Последний платеж ФЛ_28.2'}
    statuses = {'same': 'Совпадает', 'different_definition': 'Справочно: разные даты',
                'date_shift_candidate': 'Кандидат: другая дата',
                'adjacent_sum_candidate': 'Кандидат: сумма двух дней',
                'review': 'Проверить сумму',
                'missing_in_crm': 'Не найдено в CRM'}
    for row in report['rows']:
        for check in row['checks']:
            sheet.append((row['loan_id'], row['contract_no'], row['uid'],
                          row['event_time'], row['source_row_num'],
                          labels.get(check['field'], check['field']),
                          check.get('payment_date'), check['source'], check['crm'],
                          statuses[check['status']],
                          '; '.join(f"#{p['id']} {p['date']} {p['amount']}"
                                    for p in check.get('candidate_payments', [])),
                          check['note'], None, None))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476A')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((12, 29, 46, 22, 13, 30, 18, 19, 19,
                                   30, 55, 85, 27, 45), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    book.save(output)
    return output.getvalue()
